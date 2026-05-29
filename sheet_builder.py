"""
경기 마음온길 (Gyeonggi Maeum-on-Gil) — 사전정리 시트 docx 생성 모듈

summarize_to_sheet() 가 만든 dict 를 받아, 정신건강 기관에 지참할
'정신건강 상담 사전정리 시트' Word(.docx) 파일을 생성한다.

설계 원칙:
  - 문서 상단에 '자기보고형 보조 문서' 성격을 명시한다.
    (의료기록·진단서·의뢰서가 아님)
  - 8번 '안전 관련 메모' 항목 아래에는 전문가 작성용 빈칸을 둔다.
    AI 는 위험도를 등급화하지 않으므로 판단란을 비워 전달한다.
  - 위기상담 연락처(1577-0199 등)를 문서 하단에 상시 표기한다.

의존: python-docx  (pip install python-docx)
"""
from datetime import date
from io import BytesIO

from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

from prompts import _items_for_user_type, CRISIS_CONTACT_LINE


_NAVY = RGBColor(0x1F, 0x38, 0x64)
_GREY = "F2F2F2"
_LIGHT = "D6E4F0"


def _set_cell_bg(cell, hex_color: str):
    """표 셀 배경색 설정 (python-docx 기본 미지원이라 XML 직접 조작)."""
    from docx.oxml.ns import nsdecls
    from docx.oxml import parse_xml
    shading = parse_xml(
        r'<w:shd {} w:fill="{}"/>'.format(nsdecls("w"), hex_color)
    )
    cell._tc.get_or_add_tcPr().append(shading)


def build_sheet_docx(sheet_data: dict, user_type: str,
                     author_initial: str = "") -> BytesIO:
    """
    사전정리 시트 dict 를 .docx 로 변환한다.

    Parameters
    ----------
    sheet_data : dict
        summarize_to_sheet() 의 반환값
    user_type : str
        'self'(당사자) 또는 'family'(가족)
    author_initial : str
        작성자 이니셜 (선택)

    Returns
    -------
    BytesIO
        Streamlit st.download_button 에 바로 넘길 수 있는 메모리 버퍼
    """
    doc = Document()

    # 기본 글꼴
    style = doc.styles["Normal"]
    style.font.name = "맑은 고딕"
    style.font.size = Pt(10.5)

    # ---- 제목 ----
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("정신건강 상담 사전정리 시트")
    run.bold = True
    run.font.size = Pt(18)
    run.font.color.rgb = _NAVY

    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    who = "가족·보호자 작성" if user_type == "family" else "당사자 본인 작성"
    srun = sub.add_run(f"경기 마음온길(Gyeonggi Maeum-on-Gil)에서 정리  ·  {who}")
    srun.font.size = Pt(9)
    srun.font.color.rgb = RGBColor(0x59, 0x59, 0x59)

    # ---- 성격 안내 박스 ----
    notice = doc.add_paragraph()
    nrun = notice.add_run(
        "※ 이 시트는 의료기록·진단서·공식 의뢰서가 아니며, 이용자가 자신의 "
        "상태를 정리하여 상담 기관에 전달하기 위한 자기보고형 보조 문서입니다. "
        "위험도 평가 등 전문적 판단은 기관 전문가가 수행합니다."
    )
    nrun.italic = True
    nrun.font.size = Pt(8.5)
    nrun.font.color.rgb = RGBColor(0x59, 0x59, 0x59)

    # 작성일 / 작성자
    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    mrun = meta.add_run(
        f"작성일: {date.today().isoformat()}"
        + (f"   ·   작성자(이니셜): {author_initial}" if author_initial else "")
    )
    mrun.font.size = Pt(9)

    # ---- 본문 표 ----
    items = _items_for_user_type(user_type)
    table = doc.add_table(rows=1, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"

    # 헤더
    hdr = table.rows[0].cells
    hdr[0].text = "항목"
    hdr[1].text = "내용"
    for c in hdr:
        _set_cell_bg(c, "1F3864")
        for p in c.paragraphs:
            for r in p.runs:
                r.bold = True
                r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
                r.font.size = Pt(10)

    # 항목 행
    for item in items:
        row = table.add_row().cells
        # 항목명 칸
        row[0].text = ""
        p0 = row[0].paragraphs[0]
        r0 = p0.add_run(f"{item['no']}. {item['label']}")
        r0.bold = True
        r0.font.size = Pt(9.5)
        _set_cell_bg(row[0], _GREY)

        # 내용 칸
        content = sheet_data.get(item["id"], "").strip()
        row[1].text = ""
        p1 = row[1].paragraphs[0]
        if content:
            p1.add_run(content).font.size = Pt(10)
        else:
            empty = p1.add_run("(이야기 나누지 않은 항목)")
            empty.italic = True
            empty.font.size = Pt(9)
            empty.font.color.rgb = RGBColor(0x80, 0x80, 0x80)

        # 8번 안전 관련 메모: 전문가 작성용 빈칸 추가
        if item["id"] == "safety_note":
            pe = row[1].add_paragraph()
            er = pe.add_run("─" * 28)
            er.font.color.rgb = RGBColor(0xBF, 0xBF, 0xBF)
            er.font.size = Pt(8)
            pe2 = row[1].add_paragraph()
            er2 = pe2.add_run("[ 아래는 기관 전문가 작성란 — AI는 작성하지 않음 ]")
            er2.italic = True
            er2.font.size = Pt(8.5)
            er2.font.color.rgb = RGBColor(0xC5, 0x5A, 0x11)

    # 열 너비
    for r in table.rows:
        r.cells[0].width = Cm(4.2)
        r.cells[1].width = Cm(12.3)

    # ---- 하단 위기 자원 안내 ----
    doc.add_paragraph()
    crisis = doc.add_paragraph()
    crisis.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cr = crisis.add_run("힘든 마음이 클 때는 언제든 전화해 주세요  ·  " + CRISIS_CONTACT_LINE)
    cr.bold = True
    cr.font.size = Pt(9)
    cr.font.color.rgb = _NAVY

    foot = doc.add_paragraph()
    foot.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fr = foot.add_run(
        "경기 마음온길(Gyeonggi Maeum-on-Gil)  ·  정신건강 자원 안내·사전정리 AI 챗봇"
    )
    fr.font.size = Pt(8)
    fr.font.color.rgb = RGBColor(0x80, 0x80, 0x80)

    # 메모리 버퍼로 반환
    buffer = BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer
